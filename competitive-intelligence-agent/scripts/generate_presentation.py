#!/usr/bin/env python3
"""Generate the Activity 7.1 capstone presentation with screenshots and interactive elements."""

from __future__ import annotations

import io
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# CMU-inspired palette
CMU_RED = RGBColor(196, 18, 48)
DARK = RGBColor(34, 34, 34)
MID_GRAY = RGBColor(90, 90, 90)
LIGHT_BG = RGBColor(248, 249, 250)
WHITE = RGBColor(255, 255, 255)
ACCENT_BLUE = RGBColor(0, 82, 136)
GOLD = RGBColor(200, 146, 10)
CREAM = RGBColor(250, 246, 240)
GREEN = RGBColor(93, 122, 74)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
ASSETS = PROJECT_ROOT / "docs" / "presentation_assets"
OUTPUT = PROJECT_ROOT / "docs" / "Competitive-Intelligence-Agent-Capstone-Presentation.pptx"
ALT_OUTPUT = Path(
    "/Users/viveksingh/Personal/Learning/Final Project Submissions/"
    "Competitive-Intelligence-Agent-Capstone-Presentation.pptx"
)
GITHUB_URL = "https://github.com/viveksingh246/CMU-Capstone"


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_bar(slide, title: str, subtitle: str = "") -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(1.0))
    bar.fill.solid()
    bar.fill.fore_color.rgb = CMU_RED
    bar.line.fill.background()
    tf = bar.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = WHITE
    tf.margin_left = Inches(0.45)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.45), Inches(1.08), Inches(9.1), Inches(0.32))
        stf = sub.text_frame
        stf.text = subtitle
        stf.paragraphs[0].font.size = Pt(13)
        stf.paragraphs[0].font.color.rgb = MID_GRAY
        stf.paragraphs[0].font.italic = True


def add_bullets(slide, items: list[str], left=0.45, top=1.4, width=4.5, height=5.0, font_size=15):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"▸  {item}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = DARK
        p.space_after = Pt(8)
        p.line_spacing = 1.12


def add_card(slide, left, top, width, height, title, body, fill=CREAM, title_color=CMU_RED):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(232, 220, 200)
    shape.line.width = Pt(1)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = title
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = title_color
    p2 = tf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(11)
    p2.font.color.rgb = DARK
    p2.space_before = Pt(4)
    return shape


def add_hyperlink_box(slide, left, top, width, height, text, url, fill=ACCENT_BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    shape.click_action.hyperlink.address = url
    return shape


def add_image_if_exists(slide, filename: str, left, top, width, height=None) -> bool:
    path = ASSETS / filename
    if not path.exists():
        return False
    if height:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height))
    else:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(width))
    return True


def add_image_frame(slide, filename: str, left, top, width, height, caption: str = "") -> bool:
    frame = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    frame.fill.solid()
    frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = RGBColor(232, 220, 200)
    frame.line.width = Pt(1.5)
    if add_image_if_exists(slide, filename, left + 0.06, top + 0.06, width - 0.12, height - 0.35):
        if caption:
            cap = slide.shapes.add_textbox(Inches(left + 0.08), Inches(top + height - 0.28), Inches(width - 0.16), Inches(0.25))
            ctf = cap.text_frame
            ctf.text = caption
            ctf.paragraphs[0].font.size = Pt(10)
            ctf.paragraphs[0].font.color.rgb = MID_GRAY
            ctf.paragraphs[0].font.italic = True
        return True
    return False


def make_qr_png(url: str) -> io.BytesIO:
    qr = qrcode.QRCode(version=1, box_size=6, border=2)
    qr.add_data(url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#3D2914", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    return buf


def add_flow_box(slide, left, top, w, h, text, fill=WHITE):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = ACCENT_BLUE
    box.line.width = Pt(1.5)
    tf = box.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = DARK
    p.alignment = PP_ALIGN.CENTER
    return box


def add_arrow(slide, x1, y1, x2, y2):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = MID_GRAY
    conn.line.width = Pt(1.5)


def add_section_slide(prs, title: str, subtitle: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, CMU_RED)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(3.2), Inches(10), Inches(0.06))
    accent.fill.solid()
    accent.fill.fore_color.rgb = GOLD
    accent.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(2.4), Inches(8.6), Inches(2))
    tf = tb.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(38)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = WHITE
    p = tf.add_paragraph()
    p.text = subtitle
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(255, 220, 220)
    p.space_before = Pt(12)
    return slide


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # ── Slide 1: Title ──────────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, LIGHT_BG)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10), Inches(0.2))
    accent.fill.solid()
    accent.fill.fore_color.rgb = CMU_RED
    accent.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.7), Inches(1.6), Inches(5.8), Inches(2.2))
    ttf = title.text_frame
    ttf.text = "Competitive Intelligence\nResearch Agent"
    ttf.paragraphs[0].font.size = Pt(34)
    ttf.paragraphs[0].font.bold = True
    ttf.paragraphs[0].font.color.rgb = DARK
    p2 = ttf.add_paragraph()
    p2.text = "Autonomous Multi-Agent System for Source-Backed Competitor Analysis"
    p2.font.size = Pt(16)
    p2.font.color.rgb = MID_GRAY
    p2.space_before = Pt(10)

    meta = slide.shapes.add_textbox(Inches(0.7), Inches(4.2), Inches(5.5), Inches(1.2))
    mtf = meta.text_frame
    for i, line in enumerate(["Vivek Singh", "CMU Agentic AI Capstone · September 2026"]):
        p = mtf.paragraphs[0] if i == 0 else mtf.add_paragraph()
        p.text = line
        p.font.size = Pt(14 if i else 18)
        p.font.bold = i == 0
        p.font.color.rgb = DARK if i == 0 else MID_GRAY

    add_image_if_exists(slide, "01_workspace.png", 5.8, 1.3, 3.8, 4.8)
    add_hyperlink_box(slide, 0.7, 5.6, 3.2, 0.45, "View on GitHub →", GITHUB_URL)

    # ── Slide 2: Problem ────────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "The Problem & Intended Users")
    add_card(slide, 0.45, 1.35, 2.9, 1.35, "Problem", "Manual, fragmented competitive research across tabs and spreadsheets — quickly outdated")
    add_card(slide, 3.55, 1.35, 2.9, 1.35, "Gap", "Single LLM prompts lack live data, memory, validation, and follow-up research loops")
    add_card(slide, 6.65, 1.35, 2.9, 1.35, "Goal", "Autonomously plan, research, validate, compare, and report with citations")
    add_bullets(
        slide,
        [
            "Primary: Product managers & strategy analysts",
            "Secondary: Sales enablement & executives",
            "Default demo: Snowflake · Databricks · Google BigQuery",
            "Categories: Products, Features, Pricing, AI, Partnerships, Hiring",
        ],
        left=0.45,
        top=2.95,
        width=9.1,
        font_size=16,
    )

    # ── Slide 3: Why agent ────────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Why an Agent-Based Approach?", "A standalone LLM is not enough")
    cards = [
        ("Tool Calling", "Web, news, GitHub, job search APIs"),
        ("Planning", "Category-specific tasks with completion criteria"),
        ("Memory", "SQLite history + change detection across runs"),
        ("Feedback Loop", "Re-search when evidence coverage is incomplete"),
        ("Validation", "Source credibility ranking + claim labeling"),
        ("Guardrails", "Human review when confidence < 70%"),
    ]
    for i, (t, b) in enumerate(cards):
        col, row = i % 3, i // 3
        add_card(slide, 0.45 + col * 3.1, 1.4 + row * 1.55, 2.85, 1.35, t, b, fill=LIGHT_BG if row else CREAM)

    # ── Section: System ───────────────────────────────────────────────
    add_section_slide(prs, "System Design", "Architecture · Agents · Workflow")

    # ── Slide 5: Workspace screenshot ───────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Live System: Research Workspace", "Streamlit UI — configure scope and launch analysis")
    add_image_frame(slide, "01_workspace.png", 0.35, 1.25, 6.3, 5.7, "Research Workspace with industry preset, competitors, and category selection")
    add_bullets(
        slide,
        [
            "Configure industry & competitors",
            "Select analysis categories",
            "Set research period & depth",
            "One-click Run Analysis",
            "Demo / Live Search modes",
        ],
        left=6.85,
        top=1.5,
        width=2.8,
        font_size=14,
    )

    # ── Slide 6: Architecture flow ────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "System Architecture", "LangGraph orchestrates six agents with conditional routing")
    boxes = [
        (0.5, 1.35, "Streamlit UI"),
        (0.5, 2.15, "Parse & Validate"),
        (0.5, 2.95, "Check Memory"),
        (0.5, 3.75, "Plan (ReAct)"),
        (2.7, 1.35, "Search Sources"),
        (2.7, 2.15, "RAG Index"),
        (2.7, 2.95, "Extract Facts"),
        (2.7, 3.75, "Validate"),
        (4.9, 1.35, "Completeness Loop"),
        (4.9, 2.15, "ToT + Critic"),
        (4.9, 2.95, "Safety Check"),
        (4.9, 3.75, "Human Review"),
        (7.1, 1.35, "Save Memory"),
        (7.1, 2.15, "Generate Report"),
        (7.1, 2.95, "MCP Tools"),
        (7.1, 3.75, "SQLite + ChromaDB"),
    ]
    for x, y, label in boxes:
        add_flow_box(slide, x, y, 1.9, 0.55, label)
    for y in (1.6, 2.4, 3.2):
        add_arrow(slide, 2.4, y, 2.7, y)
        add_arrow(slide, 4.6, y, 4.9, y)
    add_arrow(slide, 6.8, 1.6, 7.1, 1.6)
    note = slide.shapes.add_textbox(Inches(0.5), Inches(4.55), Inches(9), Inches(0.5))
    note.text_frame.text = "Conditional edges: completeness loop (max 3×) · validation halt · human review pause at <70% confidence"
    note.text_frame.paragraphs[0].font.size = Pt(11)
    note.text_frame.paragraphs[0].font.color.rgb = MID_GRAY
    note.text_frame.paragraphs[0].font.italic = True

    # ── Slide 7: Pipeline screenshot ──────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Agent Workflow in Action", "Real-time pipeline progress during research")
    add_image_frame(slide, "02_pipeline.png", 0.35, 1.25, 9.3, 3.2, "Horizontal pipeline: Parse → Plan → Search → RAG → Extract → Validate → Analyze")
    add_card(slide, 0.35, 4.65, 2.9, 1.0, "ReAct Loop", "reason → plan → act → observe → reflect → decide")
    add_card(slide, 3.45, 4.65, 2.9, 1.0, "Completeness", "Re-search up to 3 iterations when gaps detected")
    add_card(slide, 6.55, 4.65, 2.9, 1.0, "ToT Analysis", "4 branches · beam width 3 · Critic pruning")

    # ── Slide 8: Six agents ───────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Six-Agent Multi-Agent Architecture")
    agents = [
        ("Planner", "Research plan + queries"),
        ("Research", "Collection + RAG index"),
        ("Validation", "Evidence + multi-source"),
        ("Analysis", "ToT + SWOT + scorecard"),
        ("Report", "Executive output"),
        ("Coordinator", "Memory + orchestration"),
    ]
    for i, (name, role) in enumerate(agents):
        col, row = i % 3, i // 3
        add_card(slide, 0.45 + col * 3.1, 1.35 + row * 1.7, 2.85, 1.45, name, role, fill=CREAM, title_color=ACCENT_BLUE)

    # ── Section: Evolution ────────────────────────────────────────────
    add_section_slide(prs, "Design Evolution", "Modules 1–6 · Key Refinements")

    # ── Slide 10: Evolution timeline ──────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Design Evolution Across the Program")
    timeline = [
        ("M1 · 1.1", "LangGraph workflow, SQLite, completeness loop, MCP tools"),
        ("M2 · 2.1", "ReAct reasoning + short-term memory trace in UI"),
        ("M3 · 3.1", "ChromaDB RAG — 650-token chunks, top-6 retrieval"),
        ("M4 · 4.1", "Tree-of-Thought beam search + Critic agent"),
        ("M5 · 5.1", "Six specialized agents + coordination agent"),
        ("M6 · 6.1", "Safety guardrails, escalation, human approval gate"),
    ]
    for i, (mod, desc) in enumerate(timeline):
        y = 1.35 + i * 0.88
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(y), Inches(1.1), Inches(0.55))
        badge.fill.solid()
        badge.fill.fore_color.rgb = CMU_RED if i % 2 == 0 else ACCENT_BLUE
        badge.line.fill.background()
        btf = badge.text_frame
        btf.text = mod
        btf.paragraphs[0].font.size = Pt(10)
        btf.paragraphs[0].font.bold = True
        btf.paragraphs[0].font.color.rgb = WHITE
        btf.paragraphs[0].alignment = PP_ALIGN.CENTER
        btf.vertical_anchor = MSO_ANCHOR.MIDDLE
        desc_box = slide.shapes.add_textbox(Inches(1.7), Inches(y + 0.05), Inches(7.8), Inches(0.5))
        dtf = desc_box.text_frame
        dtf.text = desc
        dtf.paragraphs[0].font.size = Pt(14)
        dtf.paragraphs[0].font.color.rgb = DARK

    # ── Section: Demo & Results ───────────────────────────────────────
    add_section_slide(prs, "Evaluation & Results", "Metrics · Screenshots · Strengths & Limitations")

    # ── Slide 12: Results screenshot ──────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Results Dashboard", "Executive brief, scorecard, recommendations, and trace")
    add_image_frame(slide, "03_results.png", 0.35, 1.25, 9.3, 3.4, "KPI strip, phase rail, recommendations, and tabbed results view")
    add_card(slide, 0.35, 4.85, 2.2, 0.9, "Coverage", "≥90% target", fill=LIGHT_BG)
    add_card(slide, 2.75, 4.85, 2.2, 0.9, "Groundedness", "≥90% cited", fill=LIGHT_BG)
    add_card(slide, 5.15, 4.85, 2.2, 0.9, "Correctness", "≥95% verified", fill=LIGHT_BG)
    add_card(slide, 7.55, 4.85, 2.0, 0.9, "Safety", "≥95% labeled", fill=LIGHT_BG)

    # ── Slide 13: Charts ──────────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Analytics Visualizations", "Plotly scorecard and competitive radar charts")
    add_image_if_exists(slide, "06_scorecard_chart.png", 0.35, 1.3, 4.5)
    add_image_if_exists(slide, "07_radar_chart.png", 5.1, 1.3, 4.5)
    cap = slide.shapes.add_textbox(Inches(0.35), Inches(5.5), Inches(9.3), Inches(0.4))
    cap.text_frame.text = "Snowflake vs. Databricks vs. Google BigQuery — weighted competitive scorecard from demo run"
    cap.text_frame.paragraphs[0].font.size = Pt(11)
    cap.text_frame.paragraphs[0].font.color.rgb = MID_GRAY
    cap.text_frame.paragraphs[0].font.italic = True

    # ── Slide 14: Metrics screenshot ──────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Evaluation Metrics & ReAct Trace", "Automated quality monitoring per research run")
    add_image_frame(slide, "05_metrics.png", 0.35, 1.25, 5.8, 5.5, "Per-run metrics: correctness, groundedness, credibility, coverage, safety")
    add_bullets(
        slide,
        [
            "131 automated pytest tests",
            "Workflow routing & halt tests",
            "RAG, ToT, Critic, safety suites",
            "Manual review of demo reports",
            "Honest gaps: reports insufficient evidence when data is thin",
        ],
        left=6.4,
        top=1.5,
        width=3.2,
        font_size=14,
    )
    add_card(slide, 6.4, 4.5, 3.2, 1.2, "Strength", "Never invents pricing or private data", fill=RGBColor(235, 245, 230), title_color=GREEN)
    add_card(slide, 6.4, 5.85, 3.2, 1.0, "Limitation", "RAG uses local embeddings in test mode", fill=RGBColor(255, 243, 230), title_color=GOLD)

    # ── Slide 15: Human review ────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Safety & Human-in-the-Loop", "Escalation when confidence falls below 70%")
    add_image_frame(slide, "04_human_review.png", 0.35, 1.25, 6.0, 5.5, "Human review pause — Approve & Generate Report gate")
    add_bullets(
        slide,
        [
            "Input validation halts bad requests",
            "2+ sources for major claims",
            "Marketing vs. verified fact labels",
            "Conflicting sources surfaced",
            "Fallbacks: sample data, ToT analyst",
            "All runs persisted to SQLite",
        ],
        left=6.6,
        top=1.5,
        width=3.0,
        font_size=14,
    )

    # ── Slide 16: Implementation ──────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Implementation Stack")
    add_card(slide, 0.45, 1.35, 2.0, 1.1, "Orchestration", "LangGraph + LangChain")
    add_card(slide, 2.65, 1.35, 2.0, 1.1, "LLM", "OpenAI GPT-4o-mini")
    add_card(slide, 4.85, 1.35, 2.0, 1.1, "Search", "Tavily API + MCP")
    add_card(slide, 7.05, 1.35, 2.0, 1.1, "Memory", "SQLite + ChromaDB")
    add_card(slide, 0.45, 2.65, 2.0, 1.1, "UI", "Streamlit + Plotly")
    add_card(slide, 2.65, 2.65, 2.0, 1.1, "Export", "Markdown / PDF / JSON")
    add_card(slide, 4.85, 2.65, 2.0, 1.1, "Testing", "131 pytest tests")
    add_card(slide, 7.05, 2.65, 2.0, 1.1, "Protocol", "MCP ci-search + ci-memory")
    add_hyperlink_box(slide, 0.45, 4.1, 4.0, 0.5, "github.com/viveksingh246/CMU-Capstone →", GITHUB_URL)
    qr_buf = make_qr_png(GITHUB_URL)
    slide.shapes.add_picture(qr_buf, Inches(5.5), Inches(3.8), width=Inches(1.6), height=Inches(1.6))
    qr_label = slide.shapes.add_textbox(Inches(5.5), Inches(5.45), Inches(1.6), Inches(0.3))
    qr_label.text_frame.text = "Scan to open repo"
    qr_label.text_frame.paragraphs[0].font.size = Pt(9)
    qr_label.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_bullets(
        slide,
        ["make setup && make run", "make test", "sample_data/ + reports/", "docs/ checkpoint mapping"],
        left=7.3,
        top=3.9,
        width=2.3,
        font_size=13,
    )

    # ── Slide 17: Closing ─────────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)
    add_title_bar(slide, "Closing Reflection")
    add_card(slide, 0.45, 1.35, 2.9, 1.5, "Main Takeaway", "Autonomous CI needs tools, memory, loops, and guardrails — not just better prompts")
    add_card(slide, 3.55, 1.35, 2.9, 1.5, "What Worked", "LangGraph orchestration, completeness loop, human review gate")
    add_card(slide, 6.65, 1.35, 2.9, 1.5, "What I Learned", "Agent design is about reliable workflows and evidence quality")
    add_bullets(
        slide,
        [
            "Next: OpenAI embeddings for production RAG quality",
            "Next: Scheduled monitoring with email/Slack alerts",
            "Next: Cloud deployment with role-based access",
            "Future: SEC filings, multi-language, knowledge graphs",
        ],
        left=0.45,
        top=3.1,
        width=9.1,
        font_size=16,
    )

    # ── Slide 18: Thank you ───────────────────────────────────────────
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, CMU_RED)
    thanks = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(8.6), Inches(2.5))
    ttf = thanks.text_frame
    ttf.text = "Thank You"
    ttf.paragraphs[0].font.size = Pt(48)
    ttf.paragraphs[0].font.bold = True
    ttf.paragraphs[0].font.color.rgb = WHITE
    ttf.paragraphs[0].alignment = PP_ALIGN.CENTER
    for text, size in [("Questions?", 22), ("github.com/viveksingh246/CMU-Capstone", 15)]:
        p = ttf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(14)
    ttf.paragraphs[-1].font.underline = True
    slide.shapes.add_picture(make_qr_png(GITHUB_URL), Inches(4.2), Inches(4.8), width=Inches(1.5), height=Inches(1.5))

    # Speaker notes
    notes = [
        (30, "Introduce project, show workspace thumbnail, mention GitHub link."),
        (50, "Problem, users, three callout cards."),
        (55, "Six capability cards — why agent not LLM."),
        (15, "Section divider — System Design."),
        (60, "Walk through live workspace screenshot."),
        (70, "Explain architecture flow diagram left to right."),
        (55, "Show pipeline screenshot — real-time progress."),
        (50, "Six agent cards — specialization."),
        (15, "Section divider — Design Evolution."),
        (55, "Timeline from Module 1 to 6."),
        (15, "Section divider — Evaluation."),
        (60, "Results dashboard screenshot + metric targets."),
        (50, "Chart screenshots — scorecard and radar."),
        (60, "Metrics screenshot + test suite + strengths/limitations."),
        (55, "Human review screenshot + safety bullets."),
        (50, "Tech stack cards, GitHub hyperlink, QR code."),
        (40, "Takeaway, learnings, next steps."),
        (25, "Thank you, QR code, invite questions."),
    ]
    for i, (seconds, text) in enumerate(notes):
        if i < len(prs.slides):
            prs.slides[i].notes_slide.notes_text_frame.text = (
                f"[~{seconds}s] {text}\n\n"
                "Delivery tip: This deck includes live UI screenshots — reference them directly. "
                f"Click the GitHub button on slides 1 and 16 during recording."
            )

    return prs


def main() -> None:
    if not ASSETS.exists():
        print("Warning: presentation_assets/ not found. Run scripts/capture_screenshots.py first.")
    prs = build_presentation()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(f"Saved: {OUTPUT}")
    try:
        ALT_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
        prs.save(ALT_OUTPUT)
        print(f"Saved: {ALT_OUTPUT}")
    except OSError as exc:
        print(f"Could not save to Final Project Submissions: {exc}")


if __name__ == "__main__":
    main()
